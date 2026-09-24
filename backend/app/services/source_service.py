"""
Source Verification and Provenance Service for RSSB Office Data Platform.
Phase 2.3 — Search Results & Source Verification.

Provides secure retrieval of source references and contextual source previews:
- Detailed SourceReference metadata without exposing internal filesystem paths
- Excel sheet/row preview with surrounding rows
- PDF page preview with bounding box coordinates and extracted text
- Human-verifiable relevance explanations
"""

import json
from pathlib import Path
from typing import Any, Dict, List, Optional
import uuid

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.models import (
    Assignment,
    Attendance,
    Document,
    DocumentVersion,
    SourceReference,
    VehicleWheelData,
)
from app.schemas.search import SourceDetailResponse, SourcePreviewResponse


def _parse_bbox(cell_or_range: Optional[str]) -> Optional[List[float]]:
    """Extracts bounding box floats from cell_or_range string if formatted as bbox:[x0,y0,x1,y1]."""
    if not cell_or_range or not cell_or_range.startswith("bbox:["):
        return None
    try:
        inner = cell_or_range[6:-1]
        coords = [float(c.strip()) for c in inner.split(",")]
        return coords if len(coords) == 4 else None
    except Exception:
        return None


def get_source_detail(db: Session, source_id: uuid.UUID) -> SourceDetailResponse:
    """
    Retrieves safe, verified metadata for a SourceReference record.
    Never exposes internal server filesystem paths.
    """
    sr = (
        db.query(SourceReference)
        .options(
            joinedload(SourceReference.document),
            joinedload(SourceReference.document_version),
        )
        .filter(SourceReference.id == source_id)
        .first()
    )

    if not sr:
        raise ValueError(f"Source reference with ID '{source_id}' was not found.")

    doc_name = sr.document.original_filename if sr.document else None
    v_num = sr.document_version.version_number if sr.document_version else None
    raw_doc_type = sr.document.document_type if sr.document else None
    doc_type = "excel" if raw_doc_type in ["xlsx", "xls"] else (raw_doc_type or "unknown")
    bbox = _parse_bbox(sr.cell_or_range)

    return SourceDetailResponse(
        id=str(sr.id),
        document_id=str(sr.document_id) if sr.document_id else None,
        document_name=doc_name,
        document_version=v_num,
        document_type=doc_type,
        page_number=sr.page_number,
        sheet_name=sr.sheet_name,
        row_number=sr.row_number,
        cell_or_range=sr.cell_or_range,
        bbox=bbox,
        source_text=sr.source_text,
        created_at=sr.created_at.isoformat() if sr.created_at else None,
    )


def get_source_preview(db: Session, source_id: uuid.UUID) -> SourcePreviewResponse:
    """
    Constructs a rich, verifiable source preview:
    - Decodes raw source text into structured key-values
    - For Excel: gathers surrounding row context if file exists, or related sheet records
    - For PDF: provides page excerpt, page number, and bounding box coordinates for highlighting
    - Explains why the source is relevant to the platform
    """
    sr = (
        db.query(SourceReference)
        .options(
            joinedload(SourceReference.document),
            joinedload(SourceReference.document_version),
        )
        .filter(SourceReference.id == source_id)
        .first()
    )

    if not sr:
        raise ValueError(f"Source reference with ID '{source_id}' was not found.")

    doc = sr.document
    version = sr.document_version
    doc_name = doc.original_filename if doc else "Unknown File"
    raw_doc_type = doc.document_type.lower() if doc and doc.document_type else "unknown"
    doc_type = "excel" if raw_doc_type in ["xlsx", "xls"] else raw_doc_type
    v_num = version.version_number if version else None

    # Parse row JSON from source_text
    parsed_content: Optional[Dict[str, Any]] = None
    if sr.source_text:
        try:
            parsed_content = json.loads(sr.source_text)
        except Exception:
            parsed_content = None

    bbox = _parse_bbox(sr.cell_or_range)

    # Find associated domain records that point to this source reference
    associated_records: List[Dict[str, Any]] = []

    # Check Attendance
    attendances = db.query(Attendance).options(joinedload(Attendance.satsang_ghar)).filter(Attendance.source_reference_id == sr.id).all()
    for att in attendances:
        associated_records.append({
            "type": "Attendance",
            "date": att.date.isoformat() if att.date else None,
            "satsang_ghar": att.satsang_ghar.name if att.satsang_ghar else None,
            "value": att.count_value,
            "raw_value": att.raw_value,
            "status": att.status,
        })

    # Check Assignment
    assignments = (
        db.query(Assignment)
        .options(
            joinedload(Assignment.satsang_ghar),
            joinedload(Assignment.person),
            joinedload(Assignment.role),
        )
        .filter(Assignment.source_reference_id == sr.id)
        .all()
    )
    for asn in assignments:
        associated_records.append({
            "type": "Assignment",
            "date": asn.date.isoformat() if asn.date else None,
            "satsang_ghar": asn.satsang_ghar.name if asn.satsang_ghar else None,
            "person": asn.person.name if asn.person else None,
            "role": asn.role.code if asn.role else None,
            "description": asn.description,
            "status": asn.status,
        })

    # Check Vehicle / Wheel
    vehicles = db.query(VehicleWheelData).options(joinedload(VehicleWheelData.satsang_ghar)).filter(VehicleWheelData.source_reference_id == sr.id).all()
    for v in vehicles:
        associated_records.append({
            "type": "Vehicle/Wheel",
            "date": v.date.isoformat() if v.date else None,
            "satsang_ghar": v.satsang_ghar.name if v.satsang_ghar else None,
            "count_value": v.count_value,
            "vehicle_type": v.vehicle_type,
            "status": v.status,
        })

    # Determine relevance explanation
    if associated_records:
        rec_summaries = []
        for r in associated_records:
            if r["type"] == "Attendance":
                rec_summaries.append(f"Attendance count {r.get('value')} for {r.get('satsang_ghar')} on {r.get('date')}")
            elif r["type"] == "Assignment":
                rec_summaries.append(f"Assignment of {r.get('person')} as {r.get('role')} at {r.get('satsang_ghar')} on {r.get('date')}")
            elif r["type"] == "Vehicle/Wheel":
                rec_summaries.append(f"{r.get('vehicle_type')} count {r.get('count_value')} for {r.get('satsang_ghar')}")
        relevance_explanation = (
            f"This source record directly provided the following domain entity in the office platform: "
            + "; ".join(rec_summaries)
            + "."
        )
    else:
        relevance_explanation = (
            f"This source reference corresponds to document '{doc_name}'"
            + (f" on sheet '{sr.sheet_name}' row {sr.row_number}" if sr.sheet_name else "")
            + (f" on page {sr.page_number}" if sr.page_number else "")
            + "."
        )

    # Build surrounding rows for Excel or page excerpt for PDF
    surrounding_rows: Optional[List[Dict[str, Any]]] = None
    page_text_excerpt: Optional[str] = None

    if doc_type in ("excel", "xlsx", "xlsm", "xls") and sr.row_number:
        target_row = sr.row_number
        if version and version.storage_path:
            storage_path = Path(version.storage_path)
            if storage_path.is_file():
                try:
                    wb = openpyxl.load_workbook(storage_path, read_only=True, data_only=True)
                    target_sheet = sr.sheet_name or wb.sheetnames[0]
                    if target_sheet in wb.sheetnames:
                        ws = wb[target_sheet]
                        start_row = max(1, target_row - 2)
                        end_row = target_row + 2

                        headers: List[str] = []
                        # Read header row (row 1)
                        header_cells = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))
                        if header_cells and header_cells[0]:
                            headers = [str(c).strip() if c is not None else f"Col {i+1}" for i, c in enumerate(header_cells[0])]

                        surrounding_rows = []
                        for row_idx, row_values in enumerate(
                            ws.iter_rows(min_row=start_row, max_row=end_row, values_only=True),
                            start=start_row,
                        ):
                            if any(v is not None for v in row_values):
                                row_dict = {
                                    (headers[i] if i < len(headers) else f"Col {i+1}"): v
                                    for i, v in enumerate(row_values)
                                    if v is not None
                                }
                                surrounding_rows.append({
                                    "row_number": row_idx,
                                    "is_target": (row_idx == target_row),
                                    "data": row_dict,
                                })
                    wb.close()
                except Exception:
                    surrounding_rows = None

        # Fallback to sibling SourceReferences from database if surrounding_rows not loaded from file
        if surrounding_rows is None and sr.document_id and sr.sheet_name:
            siblings = (
                db.query(SourceReference)
                .filter(
                    SourceReference.document_id == sr.document_id,
                    SourceReference.sheet_name == sr.sheet_name,
                    SourceReference.row_number.between(max(1, target_row - 2), target_row + 2),
                )
                .order_by(SourceReference.row_number.asc())
                .all()
            )
            if siblings:
                surrounding_rows = []
                for sib in siblings:
                    data = {}
                    if sib.source_text:
                        try:
                            data = json.loads(sib.source_text)
                        except Exception:
                            data = {"raw": sib.source_text}
                    surrounding_rows.append({
                        "row_number": sib.row_number,
                        "is_target": (sib.row_number == target_row),
                        "data": data,
                    })

    elif doc_type == "pdf":
        # Look for page-level source reference for broader page text context
        if sr.page_number:
            page_sr = (
                db.query(SourceReference)
                .filter(
                    SourceReference.document_id == sr.document_id,
                    SourceReference.page_number == sr.page_number,
                    SourceReference.row_number.is_(None),
                )
                .first()
            )
            if page_sr and page_sr.source_text:
                page_text_excerpt = page_sr.source_text
            elif sr.source_text:
                page_text_excerpt = sr.source_text

    return SourcePreviewResponse(
        id=str(sr.id),
        source_id=str(sr.id),
        document_id=str(sr.document_id) if sr.document_id else None,
        document_name=doc_name,
        document_version=v_num,
        document_type=doc_type,
        page_number=sr.page_number,
        sheet_name=sr.sheet_name,
        row_number=sr.row_number,
        cell_or_range=sr.cell_or_range,
        bbox=bbox,
        parsed_content=parsed_content,
        raw_content=sr.source_text,
        surrounding_rows=surrounding_rows,
        page_text_excerpt=page_text_excerpt,
        relevance_explanation=relevance_explanation,
        associated_records=associated_records,
    )
