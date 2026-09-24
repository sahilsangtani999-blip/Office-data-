"""
Script to generate synthetic test Excel fixtures for Phase 1.2 testing.

Creates test files under backend/tests/fixtures/:
- valid_attendance.xlsx
- valid_assignment.xlsx
- multi_sheet_office.xlsx
- ambiguous_columns.xlsx
- empty_sheet.xlsx
- unsupported_sample.txt
"""

from pathlib import Path
import openpyxl

FIXTURES_DIR = Path(__file__).parent / "fixtures"


def create_all_fixtures():
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)

    # 1. valid_attendance.xlsx
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.title = "Weekly Attendance"
    ws1.append(["Date", "Satsang Ghar", "Attendance Count", "Notes"])
    ws1.append(["2026-06-07", "Model Town", 450, "Regular Sangat"])
    ws1.append(["2026-06-14", "Karol Bagh", "approx 320", "Heavy rain"])
    ws1.append(["2026-06-21", "Rohini Sector 7", 580, "Special program"])
    wb1.save(FIXTURES_DIR / "valid_attendance.xlsx")

    # 2. valid_assignment.xlsx
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "Duty Schedule"
    ws2.append(["Date", "Satsang Ghar", "Speaker Name", "Duty Role", "Notes"])
    ws2.append(["2026-07-05", "Model Town", "Harish Kumar", "SK", "Morning discourse"])
    ws2.append(["2026-07-05", "Model Town", "Rajesh Sharma", "SR", "Bani Path"])
    ws2.append(["2026-07-12", "Karol Bagh", "Mohan Lal", "SK", "Morning discourse"])
    wb2.save(FIXTURES_DIR / "valid_assignment.xlsx")

    # 3. multi_sheet_office.xlsx
    wb3 = openpyxl.Workbook()
    ws3_1 = wb3.active
    ws3_1.title = "Attendance Summary"
    ws3_1.append(["Date", "Satsang Ghar", "Total Attendance"])
    ws3_1.append(["2026-08-02", "Pusa Road", 600])
    ws3_1.append(["2026-08-09", "Pusa Road", 650])

    ws3_2 = wb3.create_sheet(title="Vehicle Statistics")
    ws3_2.append(["Date", "Satsang Ghar", "2-Wheeler Count", "Notes"])
    ws3_2.append(["2026-08-02", "Pusa Road", 140, "Two-wheelers parked"])
    ws3_2.append(["2026-08-09", "Pusa Road", 165, "Two-wheelers parked"])
    wb3.save(FIXTURES_DIR / "multi_sheet_office.xlsx")

    # 4. ambiguous_columns.xlsx (unknown columns -> Needs Review)
    wb4 = openpyxl.Workbook()
    ws4 = wb4.active
    ws4.title = "Unrecognized Format"
    ws4.append(["Code ID", "Alpha Metric", "Beta Score", "Remarks"])
    ws4.append(["X-101", 85.5, 92.0, "Preliminary test"])
    ws4.append(["X-102", 78.0, 88.5, "Secondary test"])
    wb4.save(FIXTURES_DIR / "ambiguous_columns.xlsx")

    # 5. empty_sheet.xlsx
    wb5 = openpyxl.Workbook()
    ws5 = wb5.active
    ws5.title = "Blank"
    wb5.save(FIXTURES_DIR / "empty_sheet.xlsx")

    # 6. unsupported_sample.txt
    txt_path = FIXTURES_DIR / "unsupported_sample.txt"
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("This is a plain text file, not an Excel workbook.")

    print(f"Created all test fixtures in {FIXTURES_DIR}")


if __name__ == "__main__":
    create_all_fixtures()
