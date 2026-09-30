"""
Comprehensive test suite for Phase 3.1 — Reports Generation & Data Export Engine.
"""

from datetime import date
import os
from pathlib import Path
import unittest
import uuid

from fastapi.testclient import TestClient
import openpyxl
from sqlalchemy import text

from app.config import settings
from app.database import SessionLocal, engine
from app.main import app
from app.models import (
    Assignment,
    Attendance,
    Document,
    DocumentVersion,
    Person,
    Report,
    Role,
    SatsangGhar,
    User,
    VehicleWheelData,
)
from app.schemas.reports import ReportCreateRequest
from app.services.auth_service import seed_default_users
from app.services.report_service import EXPORTS_DIR, ReportService


def _clean_tables():
    """Wipes test data to ensure isolation."""
    with engine.begin() as conn:
        conn.execute(
            text(
                "TRUNCATE TABLE assignments, attendances, vehicle_wheel_data, "
                "source_references, document_versions, documents, data_sources, "
                "satsang_ghars, persons, roles, reports, users CASCADE;"
            )
        )


class TestReportsEngine(unittest.TestCase):
    """Test suite covering Phase 3.1 Reports service, API, exports, and RBAC."""

    @classmethod
    def setUpClass(cls):
        _clean_tables()
        cls.db = SessionLocal()
        cls.client = TestClient(app)

        # 1. Seed users
        seed_default_users(cls.db)

        # 2. Obtain valid tokens for each role
        cls.tokens = {}
        for role in ["admin", "reviewer", "uploader", "viewer"]:
            res = cls.client.post(
                "/api/v1/auth/login",
                json={"username": role, "password": f"{role}123"},
            )
            assert res.status_code == 200, f"Login failed for {role}: {res.text}"
            cls.tokens[role] = res.json()["access_token"]

        # 3. Seed domain records
        cls._seed_office_records()

    @classmethod
    def tearDownClass(cls):
        cls.db.close()
        _clean_tables()
        with SessionLocal() as db:
            seed_default_users(db)

    @classmethod
    def _seed_office_records(cls):
        # Satsang Ghars
        cls.ghar_sukhliya = SatsangGhar(name="Sukhliya", status="active")
        cls.ghar_kila_road = SatsangGhar(name="Kila Road", status="active")
        cls.db.add_all([cls.ghar_sukhliya, cls.ghar_kila_road])
        cls.db.commit()

        # Roles
        cls.role_sk = Role(name="SK", code="SK")
        cls.role_sr = Role(name="SR", code="SR")
        cls.db.add_all([cls.role_sk, cls.role_sr])
        cls.db.commit()

        # Persons
        cls.p1 = Person(name="Bhai Ram", status="active")
        cls.p2 = Person(name="Bhai Mohan", status="active")
        cls.db.add_all([cls.p1, cls.p2])
        cls.db.commit()

        # Attendance records (Oct 2026)
        att1 = Attendance(
            date=date(2026, 10, 4),
            satsang_ghar_id=cls.ghar_sukhliya.id,
            count_value=120,
            raw_value="120",
            status="active",
        )
        att2 = Attendance(
            date=date(2026, 10, 11),
            satsang_ghar_id=cls.ghar_sukhliya.id,
            count_value=140,
            raw_value="140",
            status="active",
        )
        att3 = Attendance(
            date=date(2026, 10, 18),
            satsang_ghar_id=cls.ghar_kila_road.id,
            count_value=100,
            raw_value="100",
            status="active",
        )
        cls.db.add_all([att1, att2, att3])

        # Duty assignments
        asgn1 = Assignment(
            date=date(2026, 10, 4),
            satsang_ghar_id=cls.ghar_sukhliya.id,
            person_id=cls.p1.id,
            role_id=cls.role_sk.id,
            status="active",
            description="Morning Satsang",
        )
        asgn2 = Assignment(
            date=date(2026, 10, 4),
            satsang_ghar_id=cls.ghar_sukhliya.id,
            person_id=cls.p2.id,
            role_id=cls.role_sr.id,
            status="active",
        )
        cls.db.add_all([asgn1, asgn2])

        # Vehicle records
        v1 = VehicleWheelData(
            date=date(2026, 10, 4),
            satsang_ghar_id=cls.ghar_sukhliya.id,
            vehicle_type="2-Wheeler",
            count_value=45,
            status="active",
        )
        v2 = VehicleWheelData(
            date=date(2026, 10, 4),
            satsang_ghar_id=cls.ghar_sukhliya.id,
            vehicle_type="4-Wheeler",
            count_value=15,
            status="active",
        )
        cls.db.add_all([v1, v2])
        cls.db.commit()

    # -------------------------------------------------------------------------
    # 1. Service Layer Tests
    # -------------------------------------------------------------------------

    def test_1_generate_monthly_report_service(self):
        """Service generates monthly report, computes KPIs, and produces valid files."""
        service = ReportService(self.db)
        req = ReportCreateRequest(
            name="Monthly Report - October 2026",
            report_type="monthly",
            year=2026,
            month=10,
        )

        report, result = service.generate_report(req, created_by="reviewer")

        self.assertIsNotNone(report.id)
        self.assertEqual(report.report_type, "monthly")
        self.assertEqual(report.status, "final")
        self.assertEqual(report.created_by, "reviewer")

        metrics = result["summary_metrics"]
        self.assertEqual(metrics["total_attendance"], 360)  # 120 + 140 + 100
        self.assertEqual(metrics["attendance_meetings"], 3)
        self.assertEqual(metrics["average_attendance"], 120.0)
        self.assertEqual(metrics["total_duties_assigned"], 2)
        self.assertEqual(metrics["total_vehicles_recorded"], 60)  # 45 + 15

        # Check export files exist
        xlsx_name = result["export_files"]["xlsx"]
        csv_name = result["export_files"]["csv"]
        xlsx_path = os.path.join(EXPORTS_DIR, xlsx_name)
        csv_path = os.path.join(EXPORTS_DIR, csv_name)

        self.assertTrue(os.path.exists(xlsx_path))
        self.assertTrue(os.path.exists(csv_path))

        # Inspect Excel sheets
        wb = openpyxl.load_workbook(xlsx_path)
        self.assertIn("Summary", wb.sheetnames)
        self.assertIn("Attendance", wb.sheetnames)
        self.assertIn("Duty Assignments", wb.sheetnames)
        self.assertIn("Vehicles", wb.sheetnames)

        # Inspect CSV content
        with open(csv_path, "r", encoding="utf-8") as f:
            content = f.read()
            self.assertIn("Monthly Report - October 2026", content)
            self.assertIn("360", content)

    def test_2_generate_attendance_summary_filtered_by_ghar(self):
        """Report scoped to Sukhliya filters out Kila Road records."""
        service = ReportService(self.db)
        req = ReportCreateRequest(
            report_type="attendance_summary",
            satsang_ghar="Sukhliya",
            year=2026,
            month=10,
        )

        report, result = service.generate_report(req, created_by="admin")
        metrics = result["summary_metrics"]

        # Only Sukhliya attendance: 120 + 140 = 260
        self.assertEqual(metrics["total_attendance"], 260)
        self.assertEqual(metrics["attendance_meetings"], 2)
        self.assertEqual(metrics["average_attendance"], 130.0)
        self.assertEqual(metrics["location_scope"], "Sukhliya")

    # -------------------------------------------------------------------------
    # 2. API Endpoints & RBAC Tests
    # -------------------------------------------------------------------------

    def test_3_api_generate_report_rbac(self):
        """Verifies role permissions on POST /api/v1/reports/generate."""
        payload = {
            "name": "Test RBAC Report",
            "report_type": "monthly",
            "year": 2026,
            "month": 10,
        }

        # 1. Invalid or missing token returns 401
        res_bad = self.client.post(
            "/api/v1/reports/generate",
            json=payload,
            headers={"Authorization": "Bearer bad-token-123"},
        )
        self.assertEqual(res_bad.status_code, 401)

        # 2. Viewer role returns 403 Forbidden
        headers_viewer = {"Authorization": f"Bearer {self.tokens['viewer']}"}
        res_viewer = self.client.post("/api/v1/reports/generate", json=payload, headers=headers_viewer)
        self.assertEqual(res_viewer.status_code, 403)

        # 3. Uploader role returns 403 Forbidden
        headers_uploader = {"Authorization": f"Bearer {self.tokens['uploader']}"}
        res_uploader = self.client.post("/api/v1/reports/generate", json=payload, headers=headers_uploader)
        self.assertEqual(res_uploader.status_code, 403)

        # 4. Reviewer role succeeds with 201 Created
        headers_reviewer = {"Authorization": f"Bearer {self.tokens['reviewer']}"}
        res_reviewer = self.client.post("/api/v1/reports/generate", json=payload, headers=headers_reviewer)
        self.assertEqual(res_reviewer.status_code, 201)
        data = res_reviewer.json()
        self.assertEqual(data["name"], "Test RBAC Report")
        self.assertEqual(data["created_by"], "reviewer")

        # 5. Admin role succeeds with 201 Created
        headers_admin = {"Authorization": f"Bearer {self.tokens['admin']}"}
        payload["name"] = "Admin Report"
        res_admin = self.client.post("/api/v1/reports/generate", json=payload, headers=headers_admin)
        self.assertEqual(res_admin.status_code, 201)

    def test_4_api_list_and_get_report(self):
        """All authenticated roles can list and retrieve reports."""
        service = ReportService(self.db)
        req = ReportCreateRequest(name="List Test Report", report_type="monthly", year=2026, month=10)
        report, _ = service.generate_report(req, created_by="admin")

        # Viewer can list
        headers_viewer = {"Authorization": f"Bearer {self.tokens['viewer']}"}
        res = self.client.get("/api/v1/reports", headers=headers_viewer)
        self.assertEqual(res.status_code, 200)
        reports = res.json()["reports"]
        self.assertGreaterEqual(len(reports), 1)

        # Viewer can get single report
        res_single = self.client.get(f"/api/v1/reports/{report.id}", headers=headers_viewer)
        self.assertEqual(res_single.status_code, 200)
        data = res_single.json()
        self.assertEqual(data["id"], str(report.id))
        self.assertIn("summary_metrics", data)

    def test_5_api_download_report_files(self):
        """Verifies downloading pre-generated .xlsx and .csv report files."""
        service = ReportService(self.db)
        req = ReportCreateRequest(name="Download Report", report_type="monthly", year=2026, month=10)
        report, _ = service.generate_report(req, created_by="reviewer")

        headers = {"Authorization": f"Bearer {self.tokens['viewer']}"}

        # Download XLSX
        res_xlsx = self.client.get(f"/api/v1/reports/{report.id}/download?format=xlsx", headers=headers)
        self.assertEqual(res_xlsx.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", res_xlsx.headers["content-type"])
        self.assertGreater(len(res_xlsx.content), 1000)

        # Download CSV
        res_csv = self.client.get(f"/api/v1/reports/{report.id}/download?format=csv", headers=headers)
        self.assertEqual(res_csv.status_code, 200)
        self.assertIn("text/csv", res_csv.headers["content-type"])
        self.assertIn(b"Download Report", res_csv.content)

    def test_6_api_delete_report_admin_only(self):
        """Only admin can delete reports."""
        service = ReportService(self.db)
        req = ReportCreateRequest(name="To Delete", report_type="monthly", year=2026, month=10)
        report, _ = service.generate_report(req, created_by="reviewer")

        # Reviewer cannot delete
        headers_rev = {"Authorization": f"Bearer {self.tokens['reviewer']}"}
        res_rev = self.client.delete(f"/api/v1/reports/{report.id}", headers=headers_rev)
        self.assertEqual(res_rev.status_code, 403)

        # Admin can delete
        headers_admin = {"Authorization": f"Bearer {self.tokens['admin']}"}
        res_admin = self.client.delete(f"/api/v1/reports/{report.id}", headers=headers_admin)
        self.assertEqual(res_admin.status_code, 200)
        self.assertTrue(res_admin.json()["deleted"])

        # Subsequent fetch returns 404
        res_get = self.client.get(f"/api/v1/reports/{report.id}", headers=headers_admin)
        self.assertEqual(res_get.status_code, 404)

    # -------------------------------------------------------------------------
    # 3. Search Integration
    # -------------------------------------------------------------------------

    def test_7_search_integration_for_reports(self):
        """Natural-language query for reports returns registered office reports."""
        service = ReportService(self.db)
        req = ReportCreateRequest(name="Monthly Report - October 2026", report_type="monthly", year=2026, month=10)
        service.generate_report(req, created_by="admin")

        headers = {"Authorization": f"Bearer {self.tokens['viewer']}"}
        query_payload = {"question": "Show monthly reports"}
        res = self.client.post("/api/v1/query", json=query_payload, headers=headers)

        self.assertEqual(res.status_code, 200)
        result = res.json()
        self.assertEqual(result["status"], "success")
        self.assertIn("Found", result["answer"])
        self.assertIn("report", result["answer"].lower())
        self.assertGreaterEqual(len(result["records"]), 1)
        self.assertIn("download_url", result["records"][0])


if __name__ == "__main__":
    unittest.main()
