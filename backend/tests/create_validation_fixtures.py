"""
Script to create synthetic Excel fixtures for validation testing.
"""
from pathlib import Path
import openpyxl

FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"
FIXTURES_DIR.mkdir(parents=True, exist_ok=True)


def create_excel_missing_date():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly Attendance"
    ws.append(["Date", "Satsang Ghar", "Attendance"])
    ws.append(["2026-09-06", "Model Town", 520])
    ws.append([None, "Karol Bagh", 410])  # Missing date
    ws.append(["2026-09-20", "Rohini Sector 7", 630])
    wb.save(FIXTURES_DIR / "excel_missing_date.xlsx")


def create_excel_invalid_date():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly Attendance"
    ws.append(["Date", "Satsang Ghar", "Attendance"])
    ws.append(["2026-09-06", "Model Town", 520])
    ws.append(["invalid-date-string", "Karol Bagh", 410])  # Invalid date
    ws.append(["2026-09-20", "Rohini Sector 7", 630])
    wb.save(FIXTURES_DIR / "excel_invalid_date.xlsx")


def create_excel_invalid_numeric():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly Attendance"
    ws.append(["Date", "Satsang Ghar", "Attendance"])
    ws.append(["2026-09-06", "Model Town", 520])
    ws.append(["2026-09-13", "Karol Bagh", "N/A - cancelled"])  # Invalid numeric count
    ws.append(["2026-09-20", "Rohini Sector 7", 630])
    wb.save(FIXTURES_DIR / "excel_invalid_numeric.xlsx")


def create_excel_duplicate_rows():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly Attendance"
    ws.append(["Date", "Satsang Ghar", "Attendance"])
    ws.append(["2026-09-06", "Model Town", 520])
    ws.append(["2026-09-06", "Model Town", 520])  # Duplicate logical row
    ws.append(["2026-09-13", "Karol Bagh", 410])
    wb.save(FIXTURES_DIR / "excel_duplicate_rows.xlsx")


def create_excel_unknown_columns():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly Attendance"
    ws.append(["Date", "Satsang Ghar", "Attendance", "Weather", "Remarks"])
    ws.append(["2026-09-06", "Model Town", 520, "Clear", "Normal attendance"])
    ws.append(["2026-09-13", "Karol Bagh", 410, "Rainy", "Late start"])
    wb.save(FIXTURES_DIR / "excel_unknown_columns.xlsx")


def create_clean_attendance():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly Attendance"
    ws.append(["Date", "Satsang Ghar", "Attendance"])
    ws.append(["2026-09-06", "Model Town", 520])
    ws.append(["2026-09-13", "Karol Bagh", 410])
    ws.append(["2026-09-20", "Rohini Sector 7", 630])
    wb.save(FIXTURES_DIR / "clean_attendance.xlsx")


if __name__ == "__main__":
    create_excel_missing_date()
    create_excel_invalid_date()
    create_excel_invalid_numeric()
    create_excel_duplicate_rows()
    create_excel_unknown_columns()
    create_clean_attendance()
    print("Successfully created validation Excel fixtures in:", FIXTURES_DIR)
