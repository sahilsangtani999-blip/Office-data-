"""
Generates synthetic search fixtures for Phase 2.2 testing.
Filename: backend/tests/fixtures/search_test_data.xlsx
"""

from pathlib import Path
import openpyxl

FIXTURES_DIR = Path(__file__).parent / "fixtures"


def create_search_fixtures():
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    target_path = FIXTURES_DIR / "search_test_data.xlsx"

    wb = openpyxl.Workbook()

    # Sheet 1: Sukhliya Attendance
    ws1 = wb.active
    ws1.title = "Sukhliya Attendance"
    ws1.append(["Date", "Satsang Ghar", "Attendance Count", "Notes"])
    ws1.append(["2026-09-06", "Sukhliya", 74, "Regular session"])
    ws1.append(["2026-09-13", "Sukhliya", 83, "Regular session"])
    ws1.append(["2026-09-20", "Sukhliya", 92, "Regular session"])
    ws1.append(["2026-09-27", "Sukhliya", 101, "Special Discourse"])

    # Sheet 2: Duty Schedule
    ws2 = wb.create_sheet(title="Duty Schedule")
    ws2.append(["Date", "Satsang Ghar", "Speaker Name", "Duty Role", "Notes"])
    ws2.append(["2026-09-06", "Sukhliya", "Ram Kumar", "SK", "Discourse"])
    ws2.append(["2026-09-06", "Sukhliya", "Sham Lal", "SR", "Path"])
    ws2.append(["2026-09-13", "Sukhliya", "Ram Kumar", "SK", "Discourse"])
    ws2.append(["2026-09-13", "Bicholi", "Mohan Das", "VIDEO CD", "Video CD playback"])
    ws2.append(["2026-09-20", "Bicholi", "Harish Kumar", "SK", "Discourse"])

    # Sheet 3: Bicholi Attendance
    ws3 = wb.create_sheet(title="Bicholi Attendance")
    ws3.append(["Date", "Satsang Ghar", "Attendance Count", "Notes"])
    ws3.append(["2026-09-06", "Bicholi", 60, "Regular"])
    ws3.append(["2026-09-13", "Bicholi", 70, "Regular"])

    wb.save(target_path)
    print(f"Generated search fixture at {target_path}")


if __name__ == "__main__":
    create_search_fixtures()
